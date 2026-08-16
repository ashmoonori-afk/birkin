from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from pathlib import Path
from typing import cast
from zipfile import ZipFile

from defusedxml import ElementTree
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

SCRIPT = Path(__file__).parents[2] / "script/qa/office_work_os_dogfood.py"
FORMATS = {"docx", "xlsx", "pptx", "pdf", "hwpx"}


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_real_document_dogfood_entry_point(tmp_path: Path) -> None:
    output = tmp_path / "evidence"
    run = subprocess.run(
        [sys.executable, str(SCRIPT), "--output-dir", str(output)],
        cwd=SCRIPT.parents[2],
        text=True,
        capture_output=True,
        check=False,
    )
    assert run.returncode == 0, run.stderr or run.stdout
    report = cast(dict[str, object], json.loads(run.stdout))
    formats = cast(dict[str, dict[str, object]], report["formats"])
    cleanup = cast(dict[str, object], report["cleanup_receipt"])
    assert report["ok"] is True
    assert set(formats) == FORMATS
    assert Path(cast(str, report["jail"])).resolve() == output.resolve()
    assert cleanup["artifacts_retained"] is True
    assert cleanup["temporary_paths_remaining"] == []
    assert cast(list[object], report["expected_refusals"])

    artifacts: dict[str, Path] = {}
    for format_name, evidence in formats.items():
        assert evidence["source_sha256_before"] == evidence["source_sha256_after"]
        reopened = cast(dict[str, object], evidence["reopen_validation"])
        assert reopened["ok"] is True
        assert evidence["expected_content_match"] is True
        operations = cast(dict[str, object], evidence["operations"])
        assert {"inspect", "extract", "validate", "compare"} <= set(operations)
        receipts = cast(list[dict[str, str]], evidence["artifacts"])
        for receipt in receipts:
            path = Path(receipt["path"])
            assert path.resolve().is_relative_to(output.resolve())
            assert _sha256(path) == receipt["sha256"]
        artifacts[format_name] = Path(cast(str, evidence["primary_artifact"]))

    assert any(
        "Dogfood" in paragraph.text
        for paragraph in Document(str(artifacts["docx"])).paragraphs
    )
    workbook = load_workbook(artifacts["xlsx"], read_only=True)
    value = cast(object, workbook["Evidence"]["A2"].value)
    workbook.close()
    assert value == 42
    presentation = Presentation(str(artifacts["pptx"]))
    assert len(presentation.slides) == 1
    with ZipFile(artifacts["pptx"]) as package:
        assert b"Dogfood" in package.read("ppt/slides/slide1.xml")
    assert "Dogfood" in "".join(
        page.extract_text() or "" for page in PdfReader(artifacts["pdf"], strict=True).pages
    )
    with ZipFile(artifacts["hwpx"]) as package:
        assert package.read("mimetype") == b"application/hwp+zip"
        section = package.read("Contents/section0.xml")
        _ = ElementTree.fromstring(section)
        assert b"Dogfood" in section
