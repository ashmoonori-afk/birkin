from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path
from typing import cast

import pytest
from pptx import Presentation

from birkin import approvals, store
from birkin.office.coordinator import (
    OfficeCaller,
    OfficeCoordinator,
    OfficeMutationRequest,
)
from birkin.office.errors import DocumentError, DocumentErrorCode
from birkin.tools import build_registry
from birkin.tools._types import ToolContext
from tests.office.fixture_builders import _write_package, build_hwpx_template

_W = (
    b'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    b'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"'
)
_DOCX_TYPES = (
    b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    b'<Override PartName="/word/document.xml" ContentType="application/'
    b'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
)
_P = (
    b'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    b'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
)


def _control(tag: bytes) -> bytes:
    return (
        b'<w:sdt><w:sdtPr><w:tag w:val="' + tag + b'"/></w:sdtPr><w:sdtContent>'
        b"<w:r><w:t>PLACEHOLDER</w:t></w:r></w:sdtContent></w:sdt>"
    )


def _paragraph(inner: bytes) -> bytes:
    return b"<w:p>" + inner + b"</w:p>"


def _docx(path: Path, body: bytes) -> Path:
    _write_package(path, {
        "[Content_Types].xml": _DOCX_TYPES,
        "word/document.xml": b"<w:document " + _W + b"><w:body>" + body + b"</w:body></w:document>",
    })
    return path


def _letter(path: Path) -> Path:
    """Three paragraphs; the first mixes plain text with the tagged control."""
    return _docx(path, b"".join((
        _paragraph(b"<w:r><w:t>Dear </w:t></w:r>" + _control(b"customer")),
        _paragraph(b"<w:r><w:t>Body 1</w:t></w:r>"),
        _paragraph(b"<w:r><w:t>Body 2</w:t></w:r>"),
    )))


def _hwpx(path: Path, fields: bytes) -> Path:
    _write_package(path, {
        "mimetype": b"application/hwp+zip",
        "Contents/section0.xml": (
            b'<hp:section xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
            + fields + b"</hp:section>"
        ),
    })
    return path


def _raw_deck(path: Path) -> Path:
    shapes = b"".join(
        b'<p:sp><p:nvPr><p:ph idx="%d"/></p:nvPr><a:t>PH%d</a:t></p:sp>' % (index, index)
        for index in range(2)
    )
    _write_package(path, {
        "[Content_Types].xml": (
            b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            b'<Override PartName="/ppt/presentation.xml" ContentType="application/'
            b'vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/></Types>'
        ),
        "ppt/presentation.xml": b"<p:presentation " + _P + b"/>",
        "ppt/slides/slide1.xml": b"<p:sld " + _P + b">" + shapes + b"</p:sld>",
    })
    return path


def _two_slide_deck(path: Path) -> Path:
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = "첫 제목"
    first.placeholders[1].text = "첫 본문"
    second = deck.slides.add_slide(deck.slide_layouts[1])
    second.shapes.title.text = "둘째 제목"
    second.placeholders[1].text = "둘째 본문"
    deck.save(str(path))
    return path


def _source(path: Path) -> dict[str, object]:
    return {"content_hash": hashlib.sha256(path.read_bytes()).hexdigest(), "uri": str(path)}


@pytest.fixture
def office(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    home = tmp_path / "home"
    (home / "office").mkdir(parents=True)
    (tmp_path / "caller").mkdir()
    monkeypatch.setenv("BIRKIN_HOME", str(home))
    return home / "office"


def _request(
    office: Path,
    source: Path,
    request_text: str,
    operations: list[dict[str, object]],
) -> list[dict[str, str]]:
    caller = office.parent.parent / "caller"
    approval = OfficeCoordinator(OfficeCaller(allowlist_root=caller, actor="tester")).request(
        OfficeMutationRequest(
            request_text=request_text,
            source=_source(source),
            outcome="양식 채우기",
            operations=tuple(operations),
            destination=caller / f"filled{source.suffix}",
        )
    )
    return cast("list[dict[str, str]]", approval["semantic_summaries"])


def _refused(
    office: Path,
    source: Path,
    request_text: str,
    operations: list[dict[str, object]],
) -> DocumentError:
    with pytest.raises(DocumentError) as raised:
        _ = _request(office, source, request_text, operations)
    jobs = office / "jobs"
    assert not jobs.exists() or not any(jobs.iterdir())
    return raised.value


def test_docx_field_summary_shows_the_control_text_not_its_paragraph(office: Path) -> None:
    source = _letter(office / "letter.docx")

    summaries = _request(
        office, source, "letter.docx Word 문서의 customer 필드를 채워줘",
        [{"field": "customer", "value": "홍길동"}],
    )

    assert summaries == [
        {"location": "docx field customer", "before": "PLACEHOLDER", "after": "홍길동"}
    ]


def test_hwpx_fields_are_summarized_one_per_operation(office: Path) -> None:
    source = build_hwpx_template(office / "contract.hwpx", ("customer", "date"))
    request = "contract.hwpx 한글 문서 양식의 필드를 채워줘"

    both = _request(office, source, request, [
        {"field": "customer", "value": "홍길동"},
        {"field": "date", "value": "2026-09-26"},
    ])
    one = _request(office, source, request, [{"field": "date", "value": "2026-09-26"}])

    assert both == [
        {"location": "hwpx field customer", "before": "PLACEHOLDER", "after": "홍길동"},
        {"location": "hwpx field date", "before": "PLACEHOLDER", "after": "2026-09-26"},
    ]
    assert one == [
        {"location": "hwpx field date", "before": "PLACEHOLDER", "after": "2026-09-26"}
    ]


def test_hwpx_field_fill_is_approved_and_exported(
    office: Path, tmp_path: Path
) -> None:
    # Given: a two-field HWPX template and the registered Office tools.
    source = build_hwpx_template(office / "contract.hwpx", ("customer", "date"))
    source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
    caller = tmp_path / "caller"
    destination = caller / "contract-filled.hwpx"
    registry = build_registry(
        ToolContext(cfg={}, client=None, cwd=caller, record_source="user:field-fill"),
        include={"documents"},
    )

    # When: both fields are requested and the approval is granted.
    proposed = registry.execute("office_job_request", {
        "request": "contract.hwpx 한글 문서 양식의 필드를 채워줘",
        "source": _source(source),
        "outcome": "계약서 필드 채우기",
        "operations": [
            {"field": "customer", "value": "홍길동"},
            {"field": "date", "value": "2026-09-26"},
        ],
        "destination": str(destination),
    })
    body = cast("dict[str, object]", json.loads(cast(str, proposed.content)))
    assert not proposed.is_error, body
    record = store.get_pending(cast(str, body["id"]))
    assert record is not None
    assert record["description"] == (
        "HWPX 필드 'customer' 변경: PLACEHOLDER → 홍길동\n"
        "HWPX 필드 'date' 변경: PLACEHOLDER → 2026-09-26"
    )
    result = approvals.approve(
        cast(str, body["id"]), approved_by="human:test", approved_via="test:field-fill"
    )

    # Then: the export carries both values and the source is untouched.
    assert result["ok"] is True, result
    with zipfile.ZipFile(destination) as package:
        section = package.read("Contents/section0.xml").decode("utf-8")
    assert "홍길동" in section and "2026-09-26" in section
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_sha256


@pytest.mark.parametrize(
    ("body", "field", "code"),
    [
        (
            _paragraph(_control(b"customer")),
            "missing",
            DocumentErrorCode.NODE_NOT_FOUND,
        ),
        (
            _paragraph(_control(b"customer")) + _paragraph(_control(b"customer")),
            "customer",
            DocumentErrorCode.AMBIGUOUS_LOCATOR,
        ),
        (
            _paragraph(b"<w:r><w:t>Intro</w:t></w:r>")
            + b"<w:tbl><w:tr><w:tc>" + _paragraph(_control(b"customer")) + b"</w:tc></w:tr></w:tbl>",
            "customer",
            DocumentErrorCode.UNSUPPORTED_EDIT,
        ),
    ],
)
def test_unprovable_docx_fields_are_refused_before_any_job(
    office: Path, body: bytes, field: str, code: DocumentErrorCode
) -> None:
    source = _docx(office / "form.docx", body)

    error = _refused(
        office, source, "form.docx Word 문서의 필드를 채워줘", [{"field": field, "value": "x"}]
    )

    assert error.code is code


def test_mixed_field_and_paragraph_operations_are_refused(office: Path) -> None:
    source = _letter(office / "letter.docx")

    error = _refused(office, source, "letter.docx Word 문서를 고쳐줘", [
        {"field": "customer", "value": "홍길동"},
        {"locator": {"format": "docx", "index": 2}, "value": "Body one"},
    ])

    assert error.code is DocumentErrorCode.PRECONDITION_FAILED


def test_empty_value_clears_a_field(office: Path) -> None:
    source = _letter(office / "letter.docx")

    summaries = _request(
        office, source, "letter.docx Word 문서의 customer 필드를 비워줘",
        [{"field": "customer", "value": ""}],
    )

    assert summaries == [{"location": "docx field customer", "before": "PLACEHOLDER", "after": ""}]


def test_same_hwpx_field_by_id_then_name_resolves_sequentially(office: Path) -> None:
    source = _hwpx(
        office / "named.hwpx",
        b'<hp:p id="P1"><hp:field id="f1" name="customer"><hp:t>PLACEHOLDER</hp:t></hp:field></hp:p>',
    )

    summaries = _request(office, source, "named.hwpx 한글 문서 필드를 채워줘", [
        {"field": "f1", "value": "첫 값"},
        {"field": "customer", "value": "둘째 값"},
    ])

    assert summaries == [
        {"location": "hwpx field f1", "before": "PLACEHOLDER", "after": "첫 값"},
        {"location": "hwpx field customer", "before": "첫 값", "after": "둘째 값"},
    ]


def test_pptx_placeholder_on_a_later_slide_names_that_slide(office: Path) -> None:
    source = _two_slide_deck(office / "deck.pptx")

    summaries = _request(office, source, "deck.pptx 발표 자료 본문을 바꿔줘", [
        {"locator": {"slide_part": "ppt/slides/slide2.xml", "placeholder_idx": 1}, "value": "새 본문"},
    ])

    assert summaries == [
        {"location": "pptx slide 2 placeholder 1", "before": "둘째 본문", "after": "새 본문"}
    ]


def test_pptx_slide_order_is_read_once_for_every_touched_slide(
    office: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from birkin.office import operation_targets

    # Given: a two-slide deck and a counted slide inventory.
    source = _two_slide_deck(office / "deck.pptx")
    inventory = operation_targets.presentation_inventory
    calls: list[int] = []

    def counted(parts: dict[str, bytes]) -> object:
        calls.append(len(parts))
        return inventory(parts)

    monkeypatch.setattr(operation_targets, "presentation_inventory", counted)

    # When: one request touches placeholders on both slides.
    summaries = _request(office, source, "deck.pptx 발표 자료 본문을 바꿔줘", [
        {"locator": {"slide_part": "ppt/slides/slide2.xml", "placeholder_idx": 1}, "value": "새 둘째"},
        {"locator": {"slide_part": "ppt/slides/slide1.xml", "placeholder_idx": 1}, "value": "새 첫째"},
    ])

    # Then: both slides are numbered from a single inventory pass.
    assert summaries == [
        {"location": "pptx slide 2 placeholder 1", "before": "둘째 본문", "after": "새 둘째"},
        {"location": "pptx slide 1 placeholder 1", "before": "첫 본문", "after": "새 첫째"},
    ]
    assert len(calls) == 1


def test_pptx_placeholder_before_is_its_own_text_not_the_slide(office: Path) -> None:
    source = _raw_deck(office / "raw.pptx")

    summaries = _request(
        office, source, "raw.pptx 발표 자료 제목을 바꿔줘",
        [{"placeholder_idx": 0, "value": "새 제목"}],
    )

    assert summaries == [
        {"location": "pptx slide 1 placeholder 0", "before": "PH0", "after": "새 제목"}
    ]


def test_pptx_title_without_an_index_is_refused_before_any_job(office: Path) -> None:
    source = _two_slide_deck(office / "deck.pptx")

    error = _refused(
        office, source, "deck.pptx 발표 자료 제목을 바꿔줘",
        [{"placeholder_idx": 0, "value": "새 제목"}],
    )

    assert error.code is DocumentErrorCode.NODE_NOT_FOUND
    assert store.list_pending() == []


def test_resolved_text_is_bounded_like_the_structured_preview(
    office: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from birkin.office import operation_targets

    monkeypatch.setattr(operation_targets, "MAX_PREVIEW_TEXT_BYTES", 5)
    source = build_hwpx_template(office / "contract.hwpx")

    error = _refused(
        office, source, "contract.hwpx 한글 문서 필드를 채워줘",
        [{"field": "customer", "value": "홍길동"}],
    )

    assert error.code is DocumentErrorCode.LIMIT_EXCEEDED
