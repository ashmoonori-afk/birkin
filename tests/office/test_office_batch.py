from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

from birkin import approvals, store
from birkin.office import batch
from birkin.office.coordinator import OfficeCaller


_CRASH_AFTER_THREE = r"""
import os
import sys
from birkin import store
from birkin.office import batch

approval_id = sys.argv[1]
record = store.get_pending(approval_id)
payload = record["payload"]
original = batch.execute_approved_office_job
completed = 0
def stop_after_three(*args, **kwargs):
    global completed
    receipt = original(*args, **kwargs)
    completed += 1
    if completed == 3:
        os._exit(73)
    return receipt
batch.execute_approved_office_job = stop_after_three
batch.execute(payload, approval_id=approval_id)
"""


def test_batch_keeps_partial_results_and_retries_failed_only(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("BIRKIN_HOME", str(tmp_path))
    sources = []
    items = []
    for index in range(2):
        source = tmp_path / "office" / f"source-{index}.docx"
        source.parent.mkdir(parents=True, exist_ok=True)
        source.write_bytes(f"source-{index}".encode())
        sources.append(source)
        items.append({"request": "Update Word document", "source": {"uri": str(source)}, "outcome": "update", "operations": [{"op": "x"}], "destination": str(tmp_path / f"out-{index}.docx")})
    monkeypatch.setattr(batch.OfficeCoordinator, "request", lambda self, request: {"destination": str(request.destination), "job_id": request.destination.stem})
    prepared = batch.prepare(items, OfficeCaller(tmp_path, "test"))

    def execute_one(plan, **kwargs):
        if plan["job_id"] == "out-1":
            raise RuntimeError("broken")
        return json.dumps({"job_id": plan["job_id"]})

    monkeypatch.setattr(batch, "execute_approved_office_job", execute_one)
    result = json.loads(batch.execute(prepared, approval_id="approval-1"))

    assert result["status"] == "partial_failed" and result["succeeded"] == 1 and result["failed"] == 1
    retry = batch.prepare([], OfficeCaller(tmp_path, "test"), retry_of=prepared["batch_id"])
    assert len(retry["plans"]) == 1 and retry["requests"][0]["destination"].endswith("out-1.docx")


def test_batch_checkpoints_success_before_interruption(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("BIRKIN_HOME", str(tmp_path))
    payload = {
        "batch_id": "a" * 32, "retry_of": None, "requests": [{}, {}],
        "plans": [{"job_id": "one"}, {"job_id": "two"}],
    }
    calls = 0

    def execute_one(plan, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise KeyboardInterrupt
        return json.dumps({"job_id": plan["job_id"]})

    monkeypatch.setattr(batch, "execute_approved_office_job", execute_one)
    try:
        batch.execute(payload, approval_id="approval-1")
    except KeyboardInterrupt:
        pass
    record = store._read_json(tmp_path / "office" / "batches" / f"{'a' * 32}.json", {})
    assert record["results"][0]["status"] == "succeeded"
    assert record["results"][1]["status"] == "running"


def test_batch_resume_finishes_when_all_running_jobs_have_exported_receipts(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("BIRKIN_HOME", str(tmp_path))
    batch_id = "b" * 32
    plan = {"job_id": "done", "authority_digest": "approved"}
    store._write_json(batch._path(batch_id), {
        "batch_id": batch_id, "requests": [{"destination": "done.xlsx"}], "plans": [plan],
        "results": [{"index": 0, "status": "running"}], "status": "running",
    })
    monkeypatch.setattr(batch, "approved_office_receipt", lambda _: json.dumps({
        "job_id": "done", "authority_digest": "approved", "state": "exported", "rollback": None,
    }))

    resumed = batch.prepare([], OfficeCaller(tmp_path, "test"), retry_of=batch_id)

    assert resumed["requests"] == [] and resumed["plans"] == []
    assert store._read_json(batch._path(batch_id), {})["status"] == "succeeded"
    assert json.loads(batch.execute(resumed, approval_id=None))["status"] == "succeeded"


def test_batch_real_process_restart_resumes_remaining_approved_scope(tmp_path: Path, monkeypatch) -> None:
    home = tmp_path / "home"
    caller = tmp_path / "caller"
    office_home = home / "office"
    office_home.mkdir(parents=True)
    caller.mkdir()
    monkeypatch.setenv("BIRKIN_HOME", str(home))
    source = office_home / "source.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = 1
    workbook.save(source)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    items = [
        {
            "request": "Update cell A1 in this Excel workbook",
            "source": {"uri": str(source), "content_hash": source_hash},
            "outcome": f"Set Revenue A1 to {index}",
            "operations": [{"cell": "A1", "value": index}],
            "destination": str(caller / f"out-{index}.xlsx"),
        }
        for index in range(10)
    ]
    prepared = batch.prepare(items, OfficeCaller(caller, "user:batch-test"))
    queued = approvals.propose(
        category="office_batch", title="batch", description="batch",
        payload=prepared, cfg={}, origin="user:batch-test",
    )
    approval_id = str(queued["id"])
    approvals.claim(approval_id, approved_by="human:test", approved_via="test")
    store.resolve_pending(approval_id, "executing")

    crashed = subprocess.run(
        [sys.executable, "-c", _CRASH_AFTER_THREE, approval_id],
        cwd=Path(__file__).parents[2],
        env={**os.environ, "BIRKIN_HOME": str(home), "PYTHONPATH": str(Path(__file__).parents[2])},
        check=False,
    )
    assert crashed.returncode == 73

    listed = next(item for item in batch.list_batches() if item["batch_id"] == prepared["batch_id"])
    assert listed["succeeded"] == 3 and listed["pending"] == 7 and listed["status"] == "running"

    resumed = batch.prepare([], OfficeCaller(caller, "user:batch-test"), retry_of=prepared["batch_id"])
    original = store._read_json(batch._path(prepared["batch_id"]), {})
    assert original["succeeded"] == 3 and original["pending"] == 7 and original["status"] == "running"
    assert [plan["job_id"] for plan in resumed["plans"]] == [plan["job_id"] for plan in prepared["plans"][3:]]
    assert resumed["requests"] == items[3:]

    retry = approvals.propose(
        category="office_batch", title="retry", description="retry",
        payload=resumed, cfg={}, origin="user:batch-test",
    )
    completed = approvals.approve(str(retry["id"]), approved_by="human:test", approved_via="test")
    result = store._read_json(batch._path(resumed["batch_id"]), {})
    original = store._read_json(batch._path(prepared["batch_id"]), {})
    assert completed["ok"] is True and result["succeeded"] == 7
    assert original["succeeded"] == 10 and original["pending"] == 0 and original["status"] == "succeeded"
    assert next(item for item in batch.list_batches() if item["batch_id"] == prepared["batch_id"])["status"] == "succeeded"
    assert len({plan["job_id"] for plan in prepared["plans"]}) == 10
    assert [load_workbook(caller / f"out-{index}.xlsx").active["A1"].value for index in range(10)] == list(range(10))
