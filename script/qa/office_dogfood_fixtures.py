"""Deterministic fixtures and independent reopen checks for office dogfood."""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import cast
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def artifact(path: Path) -> dict[str, str]:
    digest = sha256(path)
    return {
        "artifact_id": digest,
        "content_hash": digest,
        "media_type": "application/octet-stream",
        "uri": str(path.resolve()),
        "sensitivity": "internal",
        "acl_fingerprint": "dogfood-local",
    }


def receipt(path: Path) -> dict[str, str]:
    core = {"path": str(path.resolve()), "sha256": sha256(path)}
    encoded = json.dumps(core, sort_keys=True, separators=(",", ":")).encode()
    return {**core, "receipt_sha256": hashlib.sha256(encoded).hexdigest()}


def zip_fixture(
    source: Path, target: Path, replacement: tuple[str, bytes, bytes] | None = None
) -> None:
    with ZipFile(source) as incoming, ZipFile(target, "w") as outgoing:
        for name in incoming.namelist():
            payload = incoming.read(name)
            if replacement and name == replacement[0]:
                payload = payload.replace(replacement[1], replacement[2], 1)
            info = ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            outgoing.writestr(info, payload)


def docx_fixture(created: Path, target: Path) -> None:
    with ZipFile(created) as package:
        xml = package.read("word/document.xml")
    paragraph_match = re.search(
        rb"<w:p(?:\s[^>]*)?>(?:(?!</w:p>).)*PLACEHOLDER.*?</w:p>", xml, re.DOTALL
    )
    if paragraph_match is None:
        raise AssertionError("created DOCX has no placeholder paragraph")
    paragraph = paragraph_match.group(0)
    run_match = re.search(
        rb"<w:r(?:\s[^>]*)?>(?:(?!</w:r>).)*PLACEHOLDER.*?</w:r>", paragraph, re.DOTALL
    )
    if run_match is None:
        raise AssertionError("created DOCX placeholder has no run")
    run = run_match.group(0)
    control = (
        b'<w:sdt><w:sdtPr><w:tag w:val="customer"/></w:sdtPr><w:sdtContent>'
        + run
        + b"</w:sdtContent></w:sdt>"
    )
    zip_fixture(
        created, target, ("word/document.xml", paragraph, paragraph.replace(run, control, 1))
    )


def hwpx_fixture(path: Path) -> None:
    section = (
        b'<hp:section xmlns:hp="http://www.hancom.co.kr/hwpml/2011/section">'
        b'<hp:p id="P1"><hp:field id="customer"><hp:t>PLACEHOLDER</hp:t>'
        b"</hp:field></hp:p></hp:section>"
    )
    with ZipFile(path, "w") as package:
        mime = ZipInfo("mimetype", (1980, 1, 1, 0, 0, 0))
        mime.compress_type = ZIP_STORED
        package.writestr(mime, b"application/hwp+zip")
        part = ZipInfo("Contents/section0.xml", (1980, 1, 1, 0, 0, 0))
        part.compress_type = ZIP_DEFLATED
        package.writestr(part, section)


def extracted_text(body: dict[str, object]) -> str:
    spans = cast(list[dict[str, object]], body.get("spans", []))
    return "\n".join(str(span["text"]) for span in spans)


def reopen(path: Path, fmt: str) -> dict[str, object]:
    if fmt == "docx":
        from docx import Document

        _ = Document(str(path))
        validator = "python-docx"
    elif fmt == "xlsx":
        from openpyxl import load_workbook

        load_workbook(path, read_only=True).close()
        validator = "openpyxl"
    elif fmt == "pptx":
        from pptx import Presentation

        _ = Presentation(str(path))
        validator = "python-pptx"
    elif fmt == "pdf":
        from pypdf import PdfReader

        if not PdfReader(path, strict=True).pages:
            raise AssertionError("PDF has no pages")
        validator = "pypdf-strict"
    else:
        from defusedxml import ElementTree

        with ZipFile(path) as package:
            if package.read("mimetype") != b"application/hwp+zip":
                raise AssertionError("invalid HWPX mimetype")
            _ = ElementTree.fromstring(package.read("Contents/section0.xml"))
        validator = "zipfile+defusedxml"
    return {"ok": True, "validator": validator}
